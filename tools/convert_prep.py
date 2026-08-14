#!/usr/bin/env python3
"""
data/에베레스트_원가표.xlsx 의 「🧪 프렙 원가표」  ->  recipes/prep-*.md

원가표에는 **재료·투입량·수율만** 있고 만드는 과정이 없다.
과정은 tools/prep_methods.py 에 따로 적어 두고 여기서 합친다.

사용법:  python3 tools/convert_prep.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl 이 필요합니다:  pip install openpyxl")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from prep_methods import PREP_METHODS, PREP_ORDER  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "에베레스트_원가표.xlsx"
RECIPES_DIR = ROOT / "recipes"
SHEET = "🧪 프렙 원가표"

GROUP = "프렙"
# 프렙은 메뉴판 뒤에 붙는다. convert_xlsx.py 의 GROUP_ORDER 길이(9) 다음 자리.
GROUP_INDEX = 9

CATEGORY_ORDER = [
    "베이스 소스",
    "기본 페이스트",
    "마리네이드",
    "처트니",
    "유제품",
    "음료 베이스",
    "곁들임",
]


def clean(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_prep_sheet(ws) -> dict[str, dict]:
    """「🧪 <프렙명>」 블록마다 재료 표를 읽는다."""
    preps: dict[str, dict] = {}
    current: str | None = None

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        cells = [clean(c.value) for c in row]
        # 표가 A열이 아니라 B열부터 시작한다. 첫 내용 셀을 찾아 그 지점을 기준으로 삼는다.
        start = next((i for i, c in enumerate(cells) if c), None)
        if start is None:
            continue
        cells = cells[start:]
        first = cells[0]

        # 블록 머리:  🧪  커리 그레이비 소스 | | | 조리 후 중량: 59,904g   수율: 78%
        if first.startswith("🧪"):
            current = first.lstrip("🧪").strip()
            tail = " ".join(cells[1:])
            weight = re.search(r"조리 후 중량:\s*([\d,]+)\s*g", tail)
            yield_pct = re.search(r"수율:\s*(\d+)\s*%", tail)
            preps[current] = {
                "name": current,
                "yield_g": int(weight.group(1).replace(",", "")) if weight else 0,
                "yield_pct": int(yield_pct.group(1)) if yield_pct else 100,
                "ingredients": [],
                "total_cost": 0.0,
                "cost_per_g": 0.0,
            }
            continue

        if current is None:
            continue

        # 합계 행
        if first.startswith("합"):
            nums = [c for c in cells if re.fullmatch(r"-?\d+(\.\d+)?", c)]
            if len(nums) >= 2:
                preps[current]["cost_per_g"] = float(nums[-2])
                preps[current]["total_cost"] = float(nums[-1])
            current = None
            continue

        # 재료 행:  1 | 양파 | 60000 | 0.933 | 56000 | 비고
        if re.fullmatch(r"\d+", first) and len(cells) > 2 and cells[1]:
            amount = cells[2]
            preps[current]["ingredients"].append(
                {
                    "no": int(first),
                    "name": cells[1],
                    "amount": f"{int(float(amount)):,}g" if amount else "",
                    "note": cells[5] if len(cells) > 5 else "",
                }
            )

    return preps


def yaml_str(value: str) -> str:
    if value == "":
        return '""'
    if re.search(r'[:#\-\[\]{}&*!|>%@`"\',]|^\s|\s$', value):
        return '"' + value.replace("\\", "\\\\").replace('"', '\\"') + '"'
    return value


def md_cell(value: str) -> str:
    return value.replace("|", "\\|")


def render(prep: dict, method: dict, order: int) -> str:
    cook = method["cook_time"]
    nums = [int(n) for n in re.findall(r"\d+", cook.split("(")[0])]
    cook_min = nums[0] if nums else 0
    cook_max = nums[-1] if nums else cook_min

    serving = f"1배치 · 완성 {prep['yield_g']:,}g (수율 {prep['yield_pct']}%)"

    fm = [
        "---",
        f"id: {method['id']}",
        f"name: {yaml_str(prep['name'])}",
        f"nameEn: {yaml_str(method['name_en'])}",
        f"category: {yaml_str(method['category'])}",
        f"group: {GROUP}",
        f"order: {order}",
        f"serving: {yaml_str(serving)}",
        f"cookTime: {yaml_str(cook)}",
        f"cookTimeMin: {cook_min}",
        f"cookTimeMax: {cook_max}",
        "image: null",
        f"imageNote: {yaml_str(method['note'])}",
        f"ingredientCount: {len(prep['ingredients'])}",
        f"stepCount: {len(method['steps'])}",
        "tags:",
        "  - 프렙",
    ]
    fm.append("---")

    body = ["", "## 재료", "", "| # | 재료명 | 수량 | 비고 |", "|---:|---|---|---|"]
    for ing in prep["ingredients"]:
        body.append(
            f"| {ing['no']} | {md_cell(ing['name'])} | {md_cell(ing['amount'])} | {md_cell(ing['note'])} |"
        )

    body += ["", "## 조리 방법", ""]
    for i, step in enumerate(method["steps"], start=1):
        body.append(f"{i}. {step}")

    # 프렙에는 가니쉬가 없다. 한 줄 설명은 imageNote 로 나간다.
    body.append("")
    return "\n".join(fm) + "\n" + "\n".join(body)


def main() -> int:
    if not XLSX.exists():
        sys.exit(f"원가표를 찾을 수 없습니다: {XLSX}")

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"'{SHEET}' 시트가 없습니다. 있는 시트: {wb.sheetnames}")

    preps = parse_prep_sheet(wb[SHEET])
    print(f"원가표: {XLSX.name} — 프렙 {len(preps)}종 인식")

    for stale in RECIPES_DIR.glob("prep-*.md"):
        stale.unlink()

    warnings: list[str] = []
    written = 0

    for position, prep_name in enumerate(PREP_ORDER):
        prep = preps.get(prep_name)
        method = PREP_METHODS.get(prep_name)
        if prep is None:
            warnings.append(f"[원가표에 없음] {prep_name}")
            continue
        if method is None:
            warnings.append(f"[과정 미작성] {prep_name}")
            continue

        cat_index = CATEGORY_ORDER.index(method["category"]) if method["category"] in CATEGORY_ORDER else len(CATEGORY_ORDER)
        order = GROUP_INDEX * 1_000_000 + cat_index * 10_000 + position

        (RECIPES_DIR / f"{method['id']}.md").write_text(
            render(prep, method, order), encoding="utf-8"
        )
        written += 1

    for name in preps:
        if name not in PREP_ORDER:
            warnings.append(f"[순서 미지정 — 건너뜀] {name}")

    print(f"프렙 레시피 {written}개 -> {RECIPES_DIR.relative_to(ROOT)}/prep-*.md")

    if warnings:
        print(f"\n경고 {len(warnings)}건")
        for w in warnings:
            print("  -", w)

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
